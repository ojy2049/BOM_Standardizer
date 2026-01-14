#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
엑셀 출력 모듈
- 표준 BOM 엑셀 생성
- Table + AutoFilter 적용
- 조건부 서식
"""

import pandas as pd
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows
from typing import Optional, Tuple, List
from pathlib import Path

from config import STANDARD_OUTPUT_COLUMNS


class ExcelWriter:
    """엑셀 출력 클래스"""
    
    # 스타일 정의
    HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
    HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    SMD_FILL = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")  # 연한 초록
    DIP_FILL = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")  # 연한 주황
    UNKNOWN_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")  # 연한 노랑
    
    THIN_BORDER = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # 컬럼 너비 설정
    COLUMN_WIDTHS = {
        'NO': 6,
        '품목': 20,
        '스펙': 30,
        '수량': 8,
        '위치': 25,
        '장착방식': 12,
        '공식부품명': 40,
        '공급사': 12,
        '공급사부품번호': 25,
        '데이터시트URL': 50,
    }
    
    def __init__(self):
        self.df: Optional[pd.DataFrame] = None
        self.review_items: List[dict] = []
    
    def prepare_data(self, df: pd.DataFrame) -> pd.DataFrame:
        """
        데이터 준비 및 정리
        
        Args:
            df: 처리된 BOM 데이터프레임
            
        Returns:
            출력용 정리된 데이터프레임
        """
        # 표준 컬럼만 선택 (존재하는 컬럼만)
        out_cols = [c for c in STANDARD_OUTPUT_COLUMNS if c in df.columns]
        self.df = df[out_cols].copy()
        
        # 미확정 항목 분리
        self.review_items = []
        if '장착방식' in self.df.columns:
            unknown_mask = self.df['장착방식'] == '미확정'
            for idx, row in self.df[unknown_mask].iterrows():
                self.review_items.append(row.to_dict())
        
        return self.df
    
    def write(self, output_path: str, include_review_sheet: bool = True) -> Tuple[bool, str]:
        """
        엑셀 파일 생성
        
        Args:
            output_path: 출력 파일 경로
            include_review_sheet: 검토 시트 포함 여부
            
        Returns:
            (성공여부, 메시지)
        """
        if self.df is None or self.df.empty:
            return False, "저장할 데이터가 없습니다."
        
        try:
            # pandas로 기본 엑셀 작성
            with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
                # 메인 BOM 시트
                self.df.to_excel(writer, index=False, sheet_name='BOM')
                
                # 검토 시트 (미확정 항목)
                if include_review_sheet and self.review_items:
                    review_df = pd.DataFrame(self.review_items)
                    review_df.to_excel(writer, index=False, sheet_name='검토필요')
            
            # openpyxl로 스타일 및 테이블 추가
            self._apply_styles(output_path)
            
            msg = f"엑셀 저장 완료: {output_path}\n"
            msg += f"- 총 {len(self.df)}개 항목"
            if self.review_items:
                msg += f"\n- 검토 필요: {len(self.review_items)}개 (미확정)"
            
            return True, msg
            
        except Exception as e:
            return False, f"엑셀 저장 실패: {str(e)}"
    
    def _apply_styles(self, file_path: str):
        """엑셀 스타일 및 테이블 적용"""
        wb = load_workbook(file_path)
        
        # BOM 시트 스타일링
        if 'BOM' in wb.sheetnames:
            ws = wb['BOM']
            self._style_sheet(ws, create_table=True, table_name="BOMTable")
        
        # 검토필요 시트 스타일링
        if '검토필요' in wb.sheetnames:
            ws = wb['검토필요']
            self._style_sheet(ws, create_table=True, table_name="ReviewTable", 
                            highlight_all=True)
        
        wb.save(file_path)
    
    def _style_sheet(self, ws, create_table: bool = True, table_name: str = "DataTable",
                    highlight_all: bool = False):
        """개별 시트 스타일링"""
        if ws.max_row < 2:  # 헤더만 있으면 스킵
            return
        
        max_row = ws.max_row
        max_col = ws.max_column
        
        # 컬럼 너비 설정
        for col_idx in range(1, max_col + 1):
            col_letter = get_column_letter(col_idx)
            header_value = ws.cell(row=1, column=col_idx).value
            width = self.COLUMN_WIDTHS.get(header_value, 15)
            ws.column_dimensions[col_letter].width = width
        
        # 헤더 스타일
        for col_idx in range(1, max_col + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = self.HEADER_FONT
            cell.fill = self.HEADER_FILL
            cell.alignment = self.HEADER_ALIGNMENT
            cell.border = self.THIN_BORDER
        
        # 데이터 행 스타일
        mounting_col = None
        for col_idx in range(1, max_col + 1):
            if ws.cell(row=1, column=col_idx).value == '장착방식':
                mounting_col = col_idx
                break
        
        for row_idx in range(2, max_row + 1):
            # 장착방식에 따른 배경색
            if mounting_col:
                mounting_value = ws.cell(row=row_idx, column=mounting_col).value
                if highlight_all or mounting_value == '미확정':
                    fill = self.UNKNOWN_FILL
                elif mounting_value == 'SMD':
                    fill = self.SMD_FILL
                elif mounting_value == 'DIP':
                    fill = self.DIP_FILL
                else:
                    fill = None
            else:
                fill = None
            
            for col_idx in range(1, max_col + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.border = self.THIN_BORDER
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                
                if fill:
                    cell.fill = fill
                
                # URL 컬럼은 하이퍼링크로
                header = ws.cell(row=1, column=col_idx).value
                if header and 'URL' in header.upper() and cell.value:
                    try:
                        cell.hyperlink = cell.value
                        cell.font = Font(color="0563C1", underline="single")
                    except Exception:
                        pass
        
        # 테이블 생성
        if create_table and max_row > 1:
            table_ref = f"A1:{get_column_letter(max_col)}{max_row}"
            
            # 기존 테이블 제거 (있으면)
            if table_name in [t.displayName for t in ws.tables.values()]:
                del ws.tables[table_name]
            
            table = Table(displayName=table_name, ref=table_ref)
            style = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False
            )
            table.tableStyleInfo = style
            ws.add_table(table)
        
        # 첫 행 고정
        ws.freeze_panes = 'A2'
    
    def get_statistics(self) -> dict:
        """통계 정보 반환"""
        if self.df is None:
            return {}
        
        stats = {
            'total': len(self.df),
            'smd': 0,
            'dip': 0,
            'unknown': 0,
        }
        
        if '장착방식' in self.df.columns:
            stats['smd'] = len(self.df[self.df['장착방식'] == 'SMD'])
            stats['dip'] = len(self.df[self.df['장착방식'] == 'DIP'])
            stats['unknown'] = len(self.df[self.df['장착방식'] == '미확정'])
        
        return stats


def write_standard_bom(df: pd.DataFrame, output_path: str) -> Tuple[bool, str]:
    """
    표준 BOM 엑셀 출력 편의 함수
    
    Args:
        df: 처리된 BOM 데이터프레임
        output_path: 출력 파일 경로
        
    Returns:
        (성공여부, 메시지)
    """
    writer = ExcelWriter()
    writer.prepare_data(df)
    return writer.write(output_path)
